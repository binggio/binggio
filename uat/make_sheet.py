# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Ashish Kumar
"""Builds the UAT spreadsheet from uat/cases.py (and, after an agent run, fills in its results).

    python uat\\make_sheet.py            # writes uat\\UAT_Binggio_0.1.xlsx
"""
from __future__ import annotations
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from cases import CASES  # noqa: E402

F = 'Arial'
NAVY, BLUE, LINE = '101F36', '2357C6', 'DFE4EC'
INPUT = PatternFill('solid', fgColor='FFF7CC')      # cells a tester fills in
HEAD = PatternFill('solid', fgColor=NAVY)
AGENT = PatternFill('solid', fgColor='EEF1F5')
thin = Side(style='thin', color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

COLS = [('ID', 10), ('Area', 15), ('Priority', 9), ('Test', 34), ('Steps', 46), ('Expected result', 40),
        ('How it is run', 15), ('How the agent decides', 32), ('Agent hands it to you when', 32),
        ('Agent result', 12), ('Agent finding', 44), ('Your result', 12), ('Tester', 14), ('Date', 12), ('Notes', 36)]
COL = {name: i + 1 for i, (name, _) in enumerate(COLS)}
RESULTS = ['PASS', 'FAIL', 'BLOCKED', 'N/A']
AGENT_RESULTS = ['PASS', 'FAIL', 'AMBIGUOUS', 'NOT RUN']


def _cell(ws, r, c, v, bold=False, color=None, fill=None, size=10, align=WRAP):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name=F, size=size, bold=bold, color=color)
    if fill: x.fill = fill
    x.alignment = align
    return x


def build(path, results=None):
    results = results or {}
    wb = Workbook()

    # ------------------------------------------------------------------ How to use
    h = wb.active; h.title = 'How to use'
    h.column_dimensions['A'].width = 26; h.column_dimensions['B'].width = 96
    _cell(h, 1, 1, 'Binggio 0.1 user acceptance test', bold=True, size=16, color=NAVY)
    rows = [
        ('What this is', '%d end-to-end cases covering install, the app, the sample bank, adding data, discovery, onboarding, review, data changes, asking, safety, performance and the screens. Release when every P1 case passes.' % len(CASES)),
        ('The agent', 'python uat\\run_uat.py runs every case it can on a fresh simulated machine and fills in "Agent result" and "Agent finding" in a copy of this workbook (uat\\reports\\UAT_results_<time>.xlsx). Add --with-model for questions in your own words, --sqlserver "host=localhost;database=master;windows" for SQL Server.'),
        ('Agent result', 'PASS: checked against independent counts or exact values. FAIL: clearly wrong. AMBIGUOUS: the agent cannot decide; you decide. NOT RUN: needs a server, Playwright or a model key on this PC.'),
        ('What you fill in', 'Only the yellow columns: Your result (choose from the list), Tester, Date, Notes. Give every AMBIGUOUS case a verdict, and re-test FAIL cases after the fix.'),
        ('Your result', 'PASS, FAIL, BLOCKED (cannot be tested yet) or N/A (does not apply to this release).'),
        ('Example row', 'Your result: FAIL · Tester: Ashish · Date: 30-Oct-2026 · Notes: Refresh listed "accounts, loans" as changed but only accounts changed.'),
        ('Release rule', 'All P1 cases PASS (by the agent or by you), no FAIL left open, every AMBIGUOUS case decided. The Summary tab counts these.'),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        _cell(h, i, 1, k, bold=True); _cell(h, i, 2, v)
        h.row_dimensions[i].height = 44 if len(v) > 110 else 30
    _cell(h, 11, 1, 'Columns to fill in', bold=True); _cell(h, 11, 2, 'Yellow', fill=INPUT)
    _cell(h, 12, 1, 'Filled by the agent', bold=True); _cell(h, 12, 2, 'Grey', fill=AGENT)

    # ------------------------------------------------------------------ UAT cases
    ws = wb.create_sheet('UAT cases')
    for c, (name, width) in enumerate(COLS, start=1):
        _cell(ws, 1, c, name, bold=True, color='FFFFFF', fill=HEAD, align=Alignment(wrap_text=True, vertical='center'))
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 30
    for r, case in enumerate(CASES, start=2):
        res = results.get(case['id']) or {}
        vals = [case['id'], case['area'], case['priority'], case['title'], case['steps'], case['expected'], case['how'],
                case['decides'], case['ambiguous_when'] or '-', res.get('status'), (res.get('detail') or '')[:1000] or None, None, None, None, None]
        for c, v in enumerate(vals, start=1):
            fill = INPUT if c >= COL['Your result'] else AGENT if c in (COL['Agent result'], COL['Agent finding']) else None
            x = _cell(ws, r, c, v, fill=fill, bold=(c == 1))
            x.border = BOX
        ws.row_dimensions[r].height = max(45, 15 * (1 + case['steps'].count('\n') + len(case['steps']) // 60))
    last = len(CASES) + 1
    dv = DataValidation(type='list', formula1='"%s"' % ','.join(RESULTS), allow_blank=True)
    dv.error, dv.errorTitle = 'Choose PASS, FAIL, BLOCKED or N/A', 'Your result'
    ws.add_data_validation(dv); dv.add('%s2:%s%d' % (get_column_letter(COL['Your result']), get_column_letter(COL['Your result']), last))
    dd = DataValidation(type='date', allow_blank=True); ws.add_data_validation(dd)
    dd.add('%s2:%s%d' % (get_column_letter(COL['Date']), get_column_letter(COL['Date']), last))
    for c in range(2, last + 1): ws.cell(row=c, column=COL['Date']).number_format = 'dd-mmm-yyyy'
    colors = {'PASS': ('E3F3EF', '13846E'), 'FAIL': ('FBE7E9', 'B4232F'), 'AMBIGUOUS': ('FBF0DF', 'AE6910'),
              'NOT RUN': ('EEF0F3', '7B8798'), 'BLOCKED': ('FBF0DF', 'AE6910'), 'N/A': ('EEF0F3', '7B8798')}
    for col in (COL['Agent result'], COL['Your result']):
        rng = '%s2:%s%d' % (get_column_letter(col), get_column_letter(col), last)
        for word, (bg, fg) in colors.items():
            ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"%s"' % word],
                                                          fill=PatternFill('solid', fgColor=bg), font=Font(name=F, color=fg, bold=True)))
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(COLS)), last)

    # ------------------------------------------------------------------ Summary
    s = wb.create_sheet('Summary')
    s.column_dimensions['A'].width = 22
    for c in 'BCDEFGHI': s.column_dimensions[c].width = 12
    _cell(s, 1, 1, 'Summary', bold=True, size=14, color=NAVY)
    _cell(s, 2, 1, 'Counts update as results are entered. "Decided" = your result if given, otherwise the agent\'s PASS or FAIL.', color='5F6F84')
    heads = ['Area', 'Cases', 'P1', 'Agent PASS', 'Agent FAIL', 'Ambiguous', 'Not run', 'Your PASS', 'Your FAIL']
    for c, t in enumerate(heads, start=1):
        _cell(s, 4, c, t, bold=True, color='FFFFFF', fill=HEAD)
    rng = lambda col: "'UAT cases'!$%s$2:$%s$%d" % (get_column_letter(COL[col]), get_column_letter(COL[col]), last)
    areas = list(dict.fromkeys(c['area'] for c in CASES))
    for i, area in enumerate(areas, start=5):
        _cell(s, i, 1, area)
        a = '$A%d' % i
        f = [
            '=COUNTIF(%s,%s)' % (rng('Area'), a),
            '=COUNTIFS(%s,%s,%s,"P1")' % (rng('Area'), a, rng('Priority')),
            '=COUNTIFS(%s,%s,%s,"PASS")' % (rng('Area'), a, rng('Agent result')),
            '=COUNTIFS(%s,%s,%s,"FAIL")' % (rng('Area'), a, rng('Agent result')),
            '=COUNTIFS(%s,%s,%s,"AMBIGUOUS")' % (rng('Area'), a, rng('Agent result')),
            '=COUNTIFS(%s,%s,%s,"NOT RUN")' % (rng('Area'), a, rng('Agent result')),
            '=COUNTIFS(%s,%s,%s,"PASS")' % (rng('Area'), a, rng('Your result')),
            '=COUNTIFS(%s,%s,%s,"FAIL")' % (rng('Area'), a, rng('Your result')),
        ]
        for c, v in enumerate(f, start=2): _cell(s, i, c, v, align=Alignment(horizontal='right'))
    tot = len(areas) + 5
    _cell(s, tot, 1, 'Total', bold=True)
    for c in range(2, len(heads) + 1):
        L = get_column_letter(c)
        _cell(s, tot, c, '=SUM(%s5:%s%d)' % (L, L, tot - 1), bold=True, align=Alignment(horizontal='right'))
    r = tot + 2
    _cell(s, r, 1, 'Release check', bold=True, size=12, color=NAVY)
    checks = [
        ('P1 cases decided PASS', '=SUMPRODUCT((%s="P1")*(((%s="PASS"))+((%s="")*(%s="PASS"))>0))' % (rng('Priority'), rng('Your result'), rng('Your result'), rng('Agent result'))),
        ('P1 cases in total', '=COUNTIF(%s,"P1")' % rng('Priority')),
        ('Open failures', '=SUMPRODUCT(((%s="FAIL")*(%s<>"PASS"))+((%s="FAIL")*1)>0)' % (rng('Agent result'), rng('Your result'), rng('Your result'))),
        ('Ambiguous without your verdict', '=SUMPRODUCT((%s="AMBIGUOUS")*(%s=""))' % (rng('Agent result'), rng('Your result'))),
    ]
    for i, (k, f) in enumerate(checks, start=r + 1):
        _cell(s, i, 1, k); _cell(s, i, 2, f, bold=True, align=Alignment(horizontal='right'))
    ready = r + len(checks) + 1
    _cell(s, ready, 1, 'Ready to release', bold=True)
    _cell(s, ready, 2, '=IF(AND(B%d=B%d,B%d=0,B%d=0),"YES","NO")' % (r + 1, r + 2, r + 3, r + 4), bold=True, align=Alignment(horizontal='right'))
    s.conditional_formatting.add('B%d' % ready, CellIsRule(operator='equal', formula=['"YES"'], fill=PatternFill('solid', fgColor='E3F3EF'), font=Font(name=F, bold=True, color='13846E')))
    s.conditional_formatting.add('B%d' % ready, CellIsRule(operator='equal', formula=['"NO"'], fill=PatternFill('solid', fgColor='FBE7E9'), font=Font(name=F, bold=True, color='B4232F')))

    wb.active = 1
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == '__main__':
    print(build(HERE / 'UAT_Binggio_0.1.xlsx'))
